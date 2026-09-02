#!/usr/bin/env python3
"""
Legge il foglio "Match" del workbook TCG e genera src/data/seed.json.

Il seed e' il registro di partenza della dashboard: viene committato, quindi
l'app (e il build su GitHub Pages) non hanno bisogno ne' di Python ne' del
file .xlsx. Serve solo per ripartire dall'Excel o per aggiornarlo.

    python scripts/ingest_xlsx.py [file.xlsx]

Senza argomenti prende il .xlsx piu' recente dentro data/.
"""
from __future__ import annotations

import json
import sys
import unicodedata
from datetime import date, datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover - dipende dall'ambiente
    sys.exit("Serve openpyxl:  pip install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
USCITA = ROOT / "src" / "data" / "seed.json"

# Intestazioni attese nel foglio Match, nell'ordine delle colonne A-K.
# Le colonne L-R del workbook sono formule ricalcolate dalla dashboard.
COLONNE = {
    "#": "n",
    "Data": "data",
    "Formato": "formato",
    "Deck": "deck",
    "DeckList": "decklist",
    "Deck Avversario": "avversario",
    "1°/2°": "turno",
    "Risultato": "risultato",
    "TAG": "tag",
    "Torneo": "torneo",
    "Note": "note",
}


def scegli_workbook() -> Path:
    if len(sys.argv) > 1:
        f = Path(sys.argv[1])
        if not f.is_absolute():
            f = (Path.cwd() / f).resolve()
        return f
    candidati = [f for f in DATA.glob("*.xlsx") if not f.name.startswith("~$")]
    if not candidati:
        sys.exit(f"Nessun .xlsx in {DATA}. Copiaci il workbook TCG_Match.xlsx.")
    return max(candidati, key=lambda f: f.stat().st_mtime)


def testo(v) -> str:
    """Stringa ripulita: niente spazi doppi, niente a-capo di troppo."""
    if v is None:
        return ""
    s = str(v).replace("\r\n", "\n").strip()
    return " ".join(s.split()) if "\n" not in s else s.strip()


def normalizza(s: str) -> str:
    """Chiave per confronti tolleranti: minuscolo, senza accenti ne' spazi doppi."""
    s = unicodedata.normalize("NFKD", s.lower())
    s = "".join(c for c in s if not unicodedata.combining(c))
    return " ".join(s.split())


def iso_data(v) -> str:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = testo(v)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    return s


def turno(v) -> int | None:
    s = testo(v).replace("°", "")
    if s.startswith("1"):
        return 1
    if s.startswith("2"):
        return 2
    return None


def tag_canonici(v, visti: dict[str, str]) -> list[str]:
    """
    Divide i TAG sulla virgola e uniforma le maiuscole: nel workbook convivono
    "Bad Start" e "Bad start", che altrimenti diventerebbero due tag diversi.
    """
    fuori: list[str] = []
    for pezzo in testo(v).split(","):
        t = pezzo.strip()
        if not t:
            continue
        chiave = normalizza(t)
        canonico = visti.setdefault(chiave, t.title())
        if canonico not in fuori:
            fuori.append(canonico)
    return fuori


def main() -> None:
    sorgente = scegli_workbook()
    wb = openpyxl.load_workbook(sorgente, data_only=True)
    if "Match" not in wb.sheetnames:
        sys.exit(f'{sorgente.name}: manca il foglio "Match".')
    ws = wb["Match"]

    intestazioni = [testo(c) for c in next(ws.iter_rows(max_row=1, values_only=True))]
    indice = {campo: intestazioni.index(col) for col, campo in COLONNE.items() if col in intestazioni}
    mancanti = sorted(set(COLONNE.values()) - set(indice))
    if mancanti:
        sys.exit(f"Colonne mancanti nel foglio Match: {', '.join(mancanti)}")

    tag_visti: dict[str, str] = {}
    match: list[dict] = []
    scartati = 0

    for riga in ws.iter_rows(min_row=2, values_only=True):
        def campo(nome):
            return riga[indice[nome]] if indice[nome] < len(riga) else None

        deck = testo(campo("deck"))
        risultato = testo(campo("risultato")).upper()[:1]
        # Una riga senza deck o senza esito e' una riga vuota o incompleta.
        if not deck or risultato not in ("W", "L"):
            if any(testo(c) for c in riga[:11]):
                scartati += 1
            continue

        # L'id viene dalla posizione nel foglio, non dalla colonna "#": nel
        # workbook qualche numero e' ripetuto, e id doppi romperebbero le liste.
        match.append(
            {
                "id": f"xlsx-{len(match) + 1:03d}",
                "data": iso_data(campo("data")),
                "formato": testo(campo("formato")),
                "deck": deck,
                "decklist": testo(campo("decklist")),
                "avversario": testo(campo("avversario")),
                "turno": turno(campo("turno")),
                "risultato": risultato,
                "tag": tag_canonici(campo("tag"), tag_visti),
                "torneo": testo(campo("torneo")),
                "note": testo(campo("note")),
            }
        )

    match.sort(key=lambda m: (m["data"], m["id"]))

    def elenco(campo: str) -> list[str]:
        return sorted({m[campo] for m in match if m[campo]}, key=normalizza)

    seed = {
        "generatoDa": sorgente.name,
        "generatoIl": datetime.now().date().isoformat(),
        "liste": {
            "deck": elenco("deck"),
            "decklist": elenco("decklist"),
            "formato": elenco("formato"),
            "avversario": elenco("avversario"),
            "torneo": elenco("torneo"),
            "tag": sorted(set(tag_visti.values()), key=normalizza),
        },
        "match": match,
    }

    USCITA.parent.mkdir(parents=True, exist_ok=True)
    USCITA.write_text(json.dumps(seed, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")

    vinte = sum(1 for m in match if m["risultato"] == "W")
    print(f"{sorgente.name} -> {USCITA.relative_to(ROOT)}")
    print(
        f"  {len(match)} match, {vinte}V/{len(match) - vinte}S "
        f"({vinte / len(match):.1%}), {len(seed['liste']['deck'])} deck, "
        f"{len(seed['liste']['avversario'])} avversari, "
        f"{sum(1 for m in match if m['note'])} note"
        + (f", {scartati} righe incomplete ignorate" if scartati else "")
    )


if __name__ == "__main__":
    main()
